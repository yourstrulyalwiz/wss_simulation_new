"""Separate-entity debt schedule appendix for both existing PowerPoint exporters."""
from pptx.dml.color import RGBColor
from pptx.opc.packuri import PackURI
from pptx.util import Inches, Pt
import re


def append_borrowing_slides(prs, result, currency, scope="Selected area"):
    # The template exporter deletes/reorders slides. python-pptx chooses a new
    # part name by slide COUNT, which can collide with surviving high-numbered
    # parts. Give every appended slide a genuinely unused package address.
    part_numbers = [int(match.group(1)) for part in prs.part.package.iter_parts()
                    if (match := re.fullmatch(r"/ppt/slides/slide(\d+)\.xml", str(part.partname)))]
    next_number = max(part_numbers, default=0) + 1
    for sector in ("water_supply", "sanitation"):
        for pool in result[sector].get("scenario_borrowing_pools", []):
            if not pool.get("enabled") or not pool.get("loan_principal"):
                continue
            schedule = pool["schedule"]
            indices = [i for i, year in enumerate(schedule["years"])
                       if year >= next(y for j, y in enumerate(schedule["years"])
                                       if schedule["drawdowns"][j] > 0)]
            columns = [
                ("Year", "years"), ("Drawdown", "drawdowns"),
                ("Committed cash", "cash_committed_to_debt"),
                ("Scheduled service", "debt_service"), ("Paid service", "debt_service_paid"),
                ("Closing debt", "closing_debt"), ("Reserves", "retained_cash_reserve"),
                ("Shortfall", "debt_service_shortfall"),
            ]
            width = prs.slide_width / Inches(1)
            for start in range(0, len(indices), 12):
                rows = indices[start:start + 12]
                slide = prs.slides.add_slide(prs.slide_layouts[6])
                slide.part.partname = PackURI(f"/ppt/slides/slide{next_number}.xml")
                next_number += 1
                heading = slide.shapes.add_textbox(Inches(.4), Inches(.2), Inches(width-.8), Inches(.55))
                heading.text_frame.text = f'{scope} · {sector.replace("_", " ").title()} · {pool["entity_name"]}'
                heading.text_frame.paragraphs[0].font.size = Pt(18)
                note = slide.shapes.add_textbox(Inches(.4), Inches(.8), Inches(width-.8), Inches(1.55))
                note.text_frame.word_wrap = True
                note.text_frame.text = (
                    f'One separate loan. Principal: {pool["loan_principal"]:,.3f} {currency} M. '
                    f'Maturity: {pool["loan_end_year"]}. Outstanding at projection end '
                    f'({pool["projection_end_year"]}): {pool["outstanding_at_projection_end"]:,.3f} {currency} M.\n'
                    f'{pool["capacity_label"]}\n'
                    f'{pool["post_horizon_assumption"]}\n'
                    f'Reserves do not finance investment until released at maturity after obligations. '
                    f'Amounts below: real {currency} M. Interest basis: {pool["rate_basis"]}. '
                    f'Contractual payments do not shrink when cash is short.')
                for paragraph in note.text_frame.paragraphs:
                    paragraph.font.size = Pt(10)
                table = slide.shapes.add_table(
                    len(rows)+1, len(columns), Inches(.4), Inches(2.45),
                    Inches(width-.8), Inches(4.2)).table
                for c, (label, key) in enumerate(columns):
                    table.cell(0, c).text = label
                    table.cell(0, c).fill.solid()
                    table.cell(0, c).fill.fore_color.rgb = RGBColor.from_string("1E3A5F")
                    for r, i in enumerate(rows, 1):
                        table.cell(r, c).text = (str(schedule[key][i]) if key == "years"
                                                 else f'{schedule[key][i]:,.3f}')
                for r in range(len(rows)+1):
                    for c in range(len(columns)):
                        for paragraph in table.cell(r, c).text_frame.paragraphs:
                            paragraph.font.size = Pt(10)
                            if r == 0:
                                paragraph.font.color.rgb = RGBColor.from_string("FFFFFF")