"""Complete annual reconciliation appendix from the shared reporting definitions."""
from pptx.opc.packuri import PackURI
from pptx.util import Inches, Pt
from reporting import annual_reporting_tables, assumption_rows
import re


def append_reporting_slides(prs, result, inputs, scope='Selected area'):
    parts = [int(m.group(1)) for part in prs.part.package.iter_parts()
             if (m := re.fullmatch(r'/ppt/slides/slide(\d+)\.xml', str(part.partname)))]
    next_number = max(parts, default=0) + 1
    width, height = prs.slide_width / Inches(1), prs.slide_height / Inches(1)

    def page(title, headers, rows, *, text=False):
        nonlocal next_number
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.part.partname = PackURI(f'/ppt/slides/slide{next_number}.xml')
        next_number += 1
        box = slide.shapes.add_textbox(Inches(.35), Inches(.2), Inches(width-.7), Inches(.65))
        box.text_frame.text = title
        box.text_frame.word_wrap = True
        box.text_frame.paragraphs[0].font.size = Pt(15)
        table = slide.shapes.add_table(len(rows)+1, len(headers), Inches(.35), Inches(1),
                                      Inches(width-.7), Inches(height-1.4)).table
        if text:
            table.columns[0].width = Inches(2)
            table.columns[1].width = Inches(width-2.7)
        for c, header in enumerate(headers):
            table.cell(0, c).text = str(header)
        table.rows[0].height = Inches(.7)
        for r, row in enumerate(rows, 1):
            for c, value in enumerate(row):
                table.cell(r, c).text = (str(value) if text or c == 0 else f'{value:,.6f}')
        for r in range(len(rows)+1):
            for c in range(len(headers)):
                for p in table.cell(r, c).text_frame.paragraphs:
                    p.font.size = Pt(9 if text else 10)
                    p.font.bold = r == 0

    notes = assumption_rows(inputs)
    for start in range(0, len(notes), 5):
        page(f'{scope} · Methodology and saved-scenario assumptions',
             ['Assumption', 'Explanation'], notes[start:start+5], text=True)
    for sector in ('water_supply', 'sanitation'):
        for table in annual_reporting_tables(result, inputs, sector):
            for start in range(0, len(table['rows']), 12):
                page(f'{scope} · {sector.replace("_", " ").title()} · {table["title"]}',
                     table['headers'], table['rows'][start:start+12])