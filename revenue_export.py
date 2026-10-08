"""Annual revenue diagnostics shared by spreadsheet and PowerPoint exports."""


LEGACY_FIELDS = (
    ('Billed-household equivalents', 'connection_billed_households', 'HH'),
    ('Lagged billed Basic households', 'connection_billed_basic_households', 'HH'),
    ('Lagged billed SM households', 'connection_billed_sm_households', 'HH'),
    ('Billed Basic entry flow (delivery in prior year)', 'connection_billed_basic_entry_households', 'HH'),
    ('Billed Basic transferred out (prior delivery)', 'connection_billed_basic_transfer_households', 'HH'),
    ('Billed SM transferred in (prior delivery)', 'connection_billed_sm_transfer_households', 'HH'),
    ('Reference billed-household equivalents', 'connection_reference_billed_households', 'HH'),
    ('Lagged incremental NRW-tagged billed equivalents', 'nrw_tagged_billed_households', 'HH'),
    ('Physical NRW-origin households', 'nrw_origin_households', 'M HH'),
    ('Annual operating cost per billed household', 'connection_annual_cost_per_household', '{currency}/HH/year, real'),
    ('Equivalent operating cost per volume', 'connection_equivalent_marginal_cost', '{currency}/m³, real'),
    ('Household billed volume', 'household_billed_volume_million_m3', 'M m³'),
    ('Non-household billed volume', 'nonhousehold_billed_volume_million_m3', 'M m³'),
    ('Total billed volume', 'billed_volume_million_m3', 'M m³'),
    ('Reference billed volume', 'reference_billed_volume_million_m3', 'M m³'),
    ('Applicable tariff', 'applicable_tariff', '{currency}/m³, real'),
    ('Applicable collection ratio', 'applicable_collection_ratio', '0–1'),
    ('Reference collected revenue', 'reference_collected_revenue', '{currency} M'),
    ('Collected revenue', 'collected_revenue', '{currency} M'),
    ('Connection gross revenue difference', 'connection_revenue_delta', '{currency} M'),
    ('Incremental variable operating cost', 'incremental_variable_operating_cost', '{currency} M'),
    ('Revenue from new connections', 'connection_net_cash', '{currency} M'),
    ('Collection-efficiency cash', 'collection_cash', '{currency} M'),
    ('Tariff-reform cash', 'tariff_cash', '{currency} M'),
    ('Total additional net cash', 'additional_net_cash', '{currency} M'),
    ('Raw connection/exogenous billed volume', 'raw_billed_volume_million_m3', 'M m³'),
    ('Non-NRW reconciled volume', 'non_nrw_billed_volume_million_m3', 'M m³'),
    ('NRW sales volume', 'nrw_sales_volume', 'M m³'),
    ('Tagged household overlap', 'nrw_overlap_volume', 'M m³'),
    ('Physical recovery', 'nrw_physical_recovery', 'M m³'),
    ('Commercial recovery', 'nrw_commercial_recovery', 'M m³'),
    ('Unsold residual recovery', 'nrw_residual_recovery', 'M m³'),
    ('NRW gross sales cash', 'nrw_sales_cash', '{currency} M'),
    ('NRW attributed operating cost', 'nrw_operating_cost', '{currency} M'),
    ('NRW implementation cost', 'nrw_implementation_cost', '{currency} M'),
    ('NRW avoided-production-cost savings', 'nrw_avoided_cost_cash', '{currency} M'),
    ('NRW signed net cash', 'nrw_net', '{currency} M'),
    ('Eligible linked sanitation net cash', 'eligible_nrw_link_cash', '{currency} M'),
    ('Potential NRW service capacity', 'nrw_potential_upgrade_hh', 'M HH'),
    ('Delivered NRW upgrades', 'nrw_delivered_upgrade_hh', 'M HH'),
    ('Funded SM upgrades', 'funded_sm_upgrade_hh', 'M HH'),
    ('Funded Basic entries', 'funded_basic_entry_hh', 'M HH'),
    ('Microfinance delivery', 'mf_flow_hh', 'M HH'),
    ('Grant-assisted delivery', 'grant_flow_hh', 'M HH'),
    ('New cohort offers', 'microfinance_cohort_offers', 'M HH'),
    ('Offered cohort unserved', 'microfinance_cohort_unserved', 'M HH'),
    ('Self-finance cohort excluded', 'microfinance_cohort_self_excluded', 'M HH'),
    ('Closing cohort unoffered', 'microfinance_cohort_unoffered', 'M HH'),
    ('Remaining opening Basic eligibility', 'eligible_basic_remaining_hh', 'M HH'),
    ('Remaining opening lower-service eligibility', 'eligible_lower_remaining_hh', 'M HH'),
    ('NRW capacity committed', 'nrw_capacity_committed', 'M m³'),
    ('NRW capacity uncommitted', 'nrw_capacity_uncommitted', 'M m³'),
    ('Unallocated positive capital', 'unallocated_positive_capital', '{currency} M'),
    ('SM target overachievement', 'target_sm_overachievement_hh', 'M HH'),
    ('Basic-or-better target overachievement', 'target_basic_or_better_overachievement_hh', 'M HH'),
)


_OBSOLETE = {
    'connection_billed_households', 'connection_billed_basic_households',
    'connection_billed_sm_households', 'connection_billed_basic_entry_households',
    'connection_billed_basic_transfer_households', 'connection_billed_sm_transfer_households',
    'connection_reference_billed_households', 'nrw_tagged_billed_households',
    'connection_annual_cost_per_household', 'connection_equivalent_marginal_cost',
    'household_billed_volume_million_m3', 'nonhousehold_billed_volume_million_m3',
    'incremental_variable_operating_cost', 'nrw_operating_cost', 'connection_revenue_delta',
}
FIELDS = (
    ('Baseline aggregate billed volume', 'baseline_billed_volume_million_m3', 'M m³'),
    ('Connection volume before overlap', 'connection_raw_volume_million_m3', 'M m³'),
    ('Connection overlap exclusion', 'connection_overlap_volume_million_m3', 'M m³'),
    ('Connection volume after overlap', 'connection_volume_million_m3', 'M m³'),
    ('Avoided-sales exclusion', 'nrw_avoided_sales_adjustment', 'M m³'),
    ('Aggregate volume scaling proxy', 'connection_aggregate_volume_proxy', 'm³/reference served HH'),
) + tuple(field for field in LEGACY_FIELDS if field[1] not in _OBSOLETE)


def revenue_table(result, sector, currency):
    headers = ['Year', 'Pass', 'Revenue mode']
    headers += [f'{label} ({unit.format(currency=currency)})' for label, _, unit in FIELDS]
    sec = result[sector]
    rows = []
    for prefix, label in (('', 'BAU'), ('scenario_', 'Scenario')):
        status = sec.get(prefix + 'connection_revenue') or {}
        mode = 'Mixed (see area configurations)' if status.get('mixed') else (
            'Aggregate coverage expansion' if status.get('effective') else 'Exogenous')
        for i, year in enumerate(result['years']):
            values = [(sec.get(prefix + field) or [0] * len(result['years']))[i]
                      for _, field, _ in FIELDS]
            rows.append([year, label, mode, *values])
    return headers, rows


def append_revenue_slides(prs, results, currency, money_factor=1.0, display_currency=None):
    """Annual revenue appendix; retain full calibration/provenance in slide speaker notes."""
    import json
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.opc.packuri import PackURI
    display_currency = display_currency or currency
    labels = ['Year', 'Volume\n(M m³)', f'Reference\n({display_currency} M)',
              f'Collected\n({display_currency} M)', f'New connections\n({display_currency} M)',
              f'Collection cash\n({display_currency} M)', f'Tariff cash\n({display_currency} M)',
              f'NRW net\n({display_currency} M)', f'Linked net\n({display_currency} M)',
              f'Total net\n({display_currency} M)']
    keys = ('billed_volume_million_m3', 'reference_collected_revenue', 'collected_revenue',
            'connection_net_cash', 'collection_cash', 'tariff_cash', 'nrw_net',
            'eligible_nrw_link_cash', 'additional_net_cash')
    for scope, result in results.items():
        for sector, title in (('water_supply', 'Water'), ('sanitation', 'Sanitation')):
            sec = result[sector]
            status = sec.get('connection_revenue') or {}
            if not (status.get('requested') or status.get('mixed') or
                    any(sec.get('scenario_nrw_sales_volume') or []) or
                    any(sec.get('scenario_nrw_implementation_cost') or [])):
                continue
            forecast = [i for i, y in enumerate(result['years']) if y > result['end_asis_year']]
            for prefix, name in (('', 'BAU'), ('scenario_', 'Scenario')):
                for start in range(0, len(forecast), 8):
                    indexes = forecast[start:start + 8]
                    existing = {p.partname for p in prs.part.package.iter_parts()}
                    slide = prs.slides.add_slide(prs.slide_layouts[6])
                    if slide.part.partname in existing:
                        j = 1
                        while PackURI(f'/ppt/slides/slide{j}.xml') in existing:
                            j += 1
                        slide.part.partname = PackURI(f'/ppt/slides/slide{j}.xml')
                    box = slide.shapes.add_textbox(Inches(.45), Inches(.25), Inches(12.4), Inches(.6))
                    box.text = f'{scope.title()} — {title} revenue details: {name}'
                    p = box.text_frame.paragraphs[0]
                    p.font.size, p.font.bold = Pt(21), True
                    p.font.color.rgb = RGBColor(1, 73, 114)
                    mode = 'Mixed; see area results' if status.get('mixed') else (
                        'Connection-based' if status.get('effective') else 'Exogenous')
                    note = slide.shapes.add_textbox(Inches(.45), Inches(.92), Inches(12.4), Inches(.65))
                    note.text = f'{mode}. Annual signed cash flows, not financing-gap reductions. One-year billing lag; reference funding is not credited again.'
                    note.text_frame.paragraphs[0].font.size = Pt(11)
                    table = slide.shapes.add_table(len(indexes) + 1, len(labels), Inches(.45),
                                                  Inches(1.75), Inches(12.4), Inches(3.8)).table
                    rows = [labels]
                    for i in indexes:
                        values = [float((sec.get(prefix + key) or [0] * len(result['years']))[i])
                                  for key in keys]
                        rows.append([str(result['years'][i]), f'{values[0]:,.6f}',
                                     *[f'{v * money_factor:,.6f}' for v in values[1:]]])
                    for ri, row in enumerate(rows):
                        for ci, val in enumerate(row):
                            cell = table.cell(ri, ci)
                            cell.text = val
                            for paragraph in cell.text_frame.paragraphs:
                                paragraph.font.size = Pt(9)
                                paragraph.font.bold = ri == 0
                    # All diagnostics and empirical notes accompany the annual cash summary.
                    headers, raw = revenue_table(result, sector, currency)
                    notes = slide.notes_slide
                    if notes.notes_text_frame is None:
                        # The branded template's notes master omits a BODY placeholder.
                        # Create one explicitly so source assumptions are not dropped.
                        from pptx.enum.shapes import PP_PLACEHOLDER
                        notes.shapes._spTree.add_placeholder(
                            notes.shapes._next_shape_id, 'Revenue assumptions',
                            PP_PLACEHOLDER.BODY, 'horz', 'full', 1)
                    notes.notes_text_frame.text = json.dumps({
                        'configuration': sec.get(prefix + 'connection_revenue'),
                        'revenue_reconciliation': sec.get(prefix + 'revenue_reconciliation'),
                        'headers_native_currency': headers,
                        'annual_rows': [r for r in raw if r[1] == name and r[0] in [result['years'][i] for i in indexes]],
                        'limitations': status.get('warnings', []),
                    }, ensure_ascii=False)
