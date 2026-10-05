"""Household-only PowerPoint appendix shared by both deck export paths."""
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.opc.packuri import PackURI


def append_service_access_slides(prs, results):
    """Append one endline diagnostic slide per scope and sector.

    Results must already contain sums of locally assessed gaps for national scope.
    Household quantities never receive the deck's currency conversion factor.
    """
    fields = [
        ('Original SM target', 'target_hh', 0),
        ('Original basic-only target', 'target_hh', 1),
        ('SM coverage', 'bau_hh', 0),
        ('Basic-only coverage', 'bau_hh', 1),
        ('SM overachievement', 'sm_overachievement', None),
        ('Effective basic-only target after SM credit', 'effective_basic_only_target', None),
        ('Basic-only shortfall after SM credit (diagnostic)', 'adjusted_basic_only_gap', None),
        ('SM access gap', 'sm_access_gap', None),
        ('At-least-basic access gap (basic-entry costing)', 'at_least_basic_access_gap', None),
        ('Outstanding SM upgrades', 'closing_outstanding_hh', 0),
        ('Outstanding lower-to-basic entries', 'closing_outstanding_hh', 1),
    ]
    for scope, result in results.items():
        for sector, title in (('water_supply', 'Water Supply'), ('sanitation', 'Sanitation')):
            sec = result[sector]
            # Templates drop unused slides without renumbering their remaining parts.
            # python-pptx chooses len(slides)+1, which can collide with a surviving part.
            existing_names = {part.partname for part in prs.part.package.iter_parts()}
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            if slide.part.partname in existing_names:
                index = 1
                while PackURI(f'/ppt/slides/slide{index}.xml') in existing_names:
                    index += 1
                slide.part.partname = PackURI(f'/ppt/slides/slide{index}.xml')
            heading = slide.shapes.add_textbox(Inches(.55), Inches(.3), Inches(12.2), Inches(.55))
            heading.text_frame.text = f'{scope.title()} — {title}: service access gaps, {result["years"][-1]}'
            heading.text_frame.paragraphs[0].font.size = Pt(21)
            heading.text_frame.paragraphs[0].font.bold = True
            heading.text_frame.paragraphs[0].font.color.rgb = RGBColor(0x01, 0x49, 0x72)
            table = slide.shapes.add_table(len(fields) + 1, 3, Inches(.55), Inches(1.1),
                                           Inches(12.2), Inches(4.8)).table
            table.columns[0].width = Inches(8.2)
            table.columns[1].width = table.columns[2].width = Inches(2)
            rows = [['Measure (million households)', 'BAU', 'Scenario']]
            for label, key, rung in fields:
                values = []
                for prefix in ('', 'scenario_'):
                    name = 'scenario_hh' if prefix and key == 'bau_hh' else prefix + key
                    series = sec[name]
                    value = series[rung][-1] if rung is not None else series[-1]
                    values.append(f'{value:,.6f}')
                rows.append([label, *values])
            for ri, row in enumerate(rows):
                for ci, value in enumerate(row):
                    cell = table.cell(ri, ci)
                    cell.text = value
                    cell.text_frame.paragraphs[0].font.size = Pt(12)
                    cell.text_frame.paragraphs[0].font.bold = ri == 0
            note = slide.shapes.add_textbox(Inches(.55), Inches(6.05), Inches(12.2), Inches(.9))
            note.text_frame.word_wrap = True
            note.text_frame.text = (
                'Basic-only coverage can decline as households upgrade to safely managed service. '
                'SM above its target counts toward the basic minimum. Basic-entry costs follow SM + basic access. '
                'Gaps are assessed per area before aggregation; surplus elsewhere does not cancel a local deficit. '
                'Zero access gaps do not erase replacement, ancillary or accumulated financial shortfalls.'
            )
            note.text_frame.paragraphs[0].font.size = Pt(12)
