import copy
import io
import json
import unittest
import numpy as np
from model.connection_revenue import prepare_connection, annual_connection_cash
from model.engine import calculate, build_context
from model.service_history import sector_history
from demo_adapter import coerce_to_engine
from test_utility_revenue import example


def configuration(**overrides):
    cfg = dict(version=1, enabled=True, billed_share_sm=1, billed_share_basic=1,
               household_volume_share=1, marginal_cost=.2, zero_cost_confirmed=False,
               alignment='estimate', funding_reference='fixed', reference_confirmed=True,
               funding_includes_reforms=False, provenance={})
    cfg.update(overrides)
    fields = ['billed_share_sm', 'billed_share_basic', 'household_volume_share', 'marginal_cost']
    if cfg.get('alignment') == 'observation':
        fields.append('baseline_volume_mld')
    for key in ('consumption_m3', 'observed_billed_households', 'nonhousehold_growth_rate'):
        if cfg.get(key) is not None:
            fields.append(key)
    if cfg.get('funding_reference') == 'series':
        fields.append('reference_series')
    cfg['provenance'] = {key: {'source_type': 'assumed', 'reference_year': 2025,
                               'note': 'Controlled acceptance fixture, not empirical data.'}
                         for key in fields}
    return cfg


def fixture():
    ctx = {'years': np.array([2025, 2026, 2027, 2028]), 'bi': 0, 'n': 4,
           'population': np.array([500, 550, 600, 650])}
    history = np.zeros((5, 4))
    history[:, 0] = np.array([60, 40, 0, 0, 0]) / 1e6
    base = dict(version=1, volume_mld=.01 * 1000 / 365, tariff=1,
                collection_ratio=.8, reference_year=2025, growth_rate=0)
    return ctx, history, base


class ConnectionRevenueTests(unittest.TestCase):
    def prepared(self, cfg=None):
        ctx, history, base = fixture()
        status, runtime = prepare_connection(cfg or configuration(), base, ctx, history)
        self.assertTrue(status['effective'], status['errors'])
        return status, runtime

    def test_calibration_upgrade_and_delivered_growth(self):
        status, runtime = self.prepared()
        self.assertAlmostEqual(status['calibration']['consumption_m3'], 100)
        self.assertAlmostEqual(status['calibration']['baseline_billed_households'], 100)
        for sm, basic, volume in [(60, 40, .01), (70, 30, .01), (60, 50, .011)]:
            row = annual_connection_cash(runtime, sm / 1e6, basic / 1e6, 1, 1, .8)
            self.assertAlmostEqual(row['billed_volume_million_m3'], volume)
        _, sm_only = self.prepared(configuration(billed_share_basic=0))
        self.assertGreater(annual_connection_cash(sm_only, 70/1e6, 30/1e6, 1, 1, .8)['connection_net_cash'], 0)

    def test_cash_identity_raw_and_million_units_all_toggle_combinations(self):
        _, runtime = self.prepared()
        for tariff, collection, expected in [(1, .8, 600), (1, .9, 1700),
                                              (1.2, .8, 2360), (1.2, .9, 3680)]:
            row = annual_connection_cash(runtime, 60/1e6, 50/1e6, 1, tariff, collection)
            self.assertAlmostEqual(row['additional_net_cash'] * 1e6, expected)
            self.assertAlmostEqual(row['additional_net_cash'],
                                   row['collected_revenue'] - row['reference_collected_revenue']
                                   - row['incremental_variable_operating_cost'])
        row = annual_connection_cash(runtime, 60/1e6, 50/1e6, 1, 1.2, .9)
        for key, expected in [('reference_collected_revenue', 8000), ('collected_revenue', 11880),
                              ('connection_revenue_delta', 800), ('collection_cash', 1100),
                              ('tariff_cash', 1980), ('incremental_variable_operating_cost', 200),
                              ('connection_net_cash', 600)]:
            self.assertAlmostEqual(row[key] * 1e6, expected)
        _, zero = self.prepared(configuration(marginal_cost=0, zero_cost_confirmed=True))
        self.assertAlmostEqual(annual_connection_cash(zero, 60/1e6, 50/1e6, 1, 1.2, .9)['additional_net_cash']*1e6, 3880)

    def test_signed_negative_cash_and_avoided_variable_cost(self):
        _, runtime = self.prepared()
        row = annual_connection_cash(runtime, 50/1e6, 40/1e6, 1, 1, .8)
        self.assertAlmostEqual(row['connection_revenue_delta']*1e6, -800)
        self.assertAlmostEqual(row['incremental_variable_operating_cost']*1e6, -200)
        self.assertAlmostEqual(row['connection_net_cash']*1e6, -600)

    def test_validation_no_silent_defaults_and_valid_zeros(self):
        ctx, history, base = fixture()
        cases = [dict(billed_share_sm=None), dict(marginal_cost=0), dict(reference_confirmed=False),
                 dict(funding_includes_reforms=True), dict(funding_reference='unknown'),
                 dict(billed_share_basic=1.1), dict(observed_billed_households=101),
                 dict(alignment='observation', baseline_volume_mld=None)]
        for override in cases:
            status, runtime = prepare_connection(configuration(**override), base, ctx, history)
            self.assertFalse(status['effective'], override)
            self.assertIsNone(runtime)
            self.assertTrue(status['errors'])
        cfg = configuration()
        cfg['provenance'] = {}
        self.assertFalse(prepare_connection(cfg, base, ctx, history)[0]['effective'])
        status, runtime = self.prepared(configuration(billed_share_sm=0, billed_share_basic=1))
        self.assertEqual(runtime['fsm'], 0)
        self.assertAlmostEqual(status['calibration']['baseline_billed_households'], 40)

    def test_zero_calibration_and_nonhousehold_only(self):
        ctx, history, base = fixture()
        history[:2, 0] = 0
        self.assertFalse(prepare_connection(configuration(), base, ctx, history)[0]['effective'])
        base['volume_mld'] = 0
        self.assertFalse(prepare_connection(configuration(), base, ctx, history)[0]['effective'])
        status, runtime = prepare_connection(configuration(consumption_m3=100), base, ctx, history)
        self.assertTrue(status['effective'])
        self.assertAlmostEqual(annual_connection_cash(runtime, 10/1e6, 0, 1, 1, .8)['billed_volume_million_m3'], .001)
        base['volume_mld'] = .01*1000/365
        status, runtime = prepare_connection(configuration(household_volume_share=0), base, ctx, history)
        self.assertTrue(status['effective'])
        self.assertEqual(annual_connection_cash(runtime, 100/1e6, 0, 1, 1, .8)['connection_net_cash'], 0)

    def test_reference_alignment_freezing_and_population_not_double_scaled(self):
        ctx, history, base = fixture()
        base.update(reference_year=2026, growth_rate=.1)
        cfg = configuration(alignment=None)
        self.assertFalse(prepare_connection(cfg, base, ctx, history)[0]['effective'])
        status, runtime = prepare_connection(configuration(), base, ctx, history)
        self.assertAlmostEqual(status['calibration']['baseline_volume_million_m3'], .01/1.1)
        self.assertEqual(status['calibration']['original_reference_year'], 2026)
        observed = configuration(alignment='observation', baseline_volume_mld=.02*1000/365)
        self.assertAlmostEqual(prepare_connection(observed, base, ctx, history)[0]['calibration']['baseline_volume_million_m3'], .02)
        base.update(reference_year=2025, growth_rate=None)
        fixed = prepare_connection(configuration(), base, ctx, history)[1]
        exog = prepare_connection(configuration(funding_reference='exogenous'), base, ctx, history)[1]
        for t in (1, 2, 3):
            row = annual_connection_cash(exog, 60/1e6, 40/1e6, t, 1, .8)
            self.assertAlmostEqual(row['household_billed_volume_million_m3'], .01)
            self.assertLess(row['connection_net_cash'], 0)
            self.assertEqual(annual_connection_cash(fixed, 60/1e6, 40/1e6, t, 1, .8)['connection_net_cash'], 0)
        series = configuration(funding_reference='series', reference_series={'2026':10000,'2027':11000,'2028':12000})
        self.assertTrue(prepare_connection(series, base, ctx, history)[0]['effective'])
        del series['reference_series']['2028']
        self.assertFalse(prepare_connection(series, base, ctx, history)[0]['effective'])

    def test_mixed_volume_path(self):
        _, runtime = self.prepared(configuration(household_volume_share=.5, nonhousehold_growth_rate=.1))
        row = annual_connection_cash(runtime, 60/1e6, 40/1e6, 1, 1, .8)
        self.assertAlmostEqual(row['household_billed_volume_million_m3'], .005)
        self.assertAlmostEqual(row['nonhousehold_billed_volume_million_m3'], .0055)
        self.assertEqual(row['connection_net_cash'], 0)

    def test_integrated_lag_capital_once_bau_independence_and_reload(self):
        d = example()
        d['connection_revenue'] = {s: configuration(billed_share_basic=0) for s in ('water','sanitation')}
        before = json.dumps(d, sort_keys=True)
        engine = coerce_to_engine(d)
        result = calculate(engine)
        self.assertEqual(json.dumps(d, sort_keys=True), before)
        by = result['years'].index(d['period']['baseline_year'])
        for sector in ('water_supply', 'sanitation'):
            sec = result[sector]
            self.assertTrue(sec['connection_revenue']['effective'])
            self.assertAlmostEqual(sec['connection_net_cash'][by+1], 0)
            q = sec['connection_revenue']['calibration']['consumption_m3']
            for t in range(by+1, len(result['years'])):
                self.assertAlmostEqual(sec['household_billed_volume_million_m3'][t], q * sec['bau_hh'][0][t-1])
                self.assertAlmostEqual(sec['available_total'][t]-sec['bau_available'][t], sec['additional_net_cash'][t])
            self.assertEqual(sec['scenario_connection_net_cash'], sec['connection_net_cash'])
        self.assertEqual(result, calculate(coerce_to_engine(json.loads(json.dumps(d)))))
        d['toggles']['ws_tariff_enabled'] = True
        on = calculate(coerce_to_engine(d))
        for key in ('bau_hh','financing_gap','connection_net_cash','funded_asset_stock'):
            self.assertEqual(on['water_supply'][key], result['water_supply'][key])
        self.assertEqual(on['sanitation']['scenario_connection_net_cash'], result['sanitation']['connection_net_cash'])

    def test_incomplete_config_falls_back_visibly_and_off_restores(self):
        d = example()
        old = calculate(coerce_to_engine(d))
        d['connection_revenue'] = {'water': {'version':1, 'enabled':True}}
        result = calculate(coerce_to_engine(d))
        self.assertTrue(result['water_supply']['connection_revenue']['errors'])
        self.assertFalse(result['water_supply']['connection_revenue']['effective'])
        for key in ('bau_hh','financing_gap','available_total','billed_volume_million_m3'):
            self.assertEqual(result['water_supply'][key], old['water_supply'][key])
        d['connection_revenue']['water'] = configuration(enabled=False)
        result = calculate(coerce_to_engine(d))
        for key in ('bau_hh','financing_gap','available_total','billed_volume_million_m3'):
            self.assertEqual(result['water_supply'][key], old['water_supply'][key])

    def test_endpoint_calibration_matches_delivery_history(self):
        from app import revenue_bases
        d = example()
        d['connection_revenue'] = {s: configuration() for s in ('water','sanitation')}
        resolved = revenue_bases(copy.deepcopy(d))
        result = calculate(coerce_to_engine(d))
        for sector, key in [('water','water_supply'),('sanitation','sanitation')]:
            self.assertEqual(resolved[sector]['connection']['calibration'], result[key]['connection_revenue']['calibration'])

    def test_cumulative_passes_use_their_own_lagged_connections(self):
        from deck_data import cumulative_passes
        d = example()
        d['connection_revenue'] = {s: configuration(billed_share_basic=0) for s in ('water','sanitation')}
        for prefix in ('ws','san'):
            for lever in ('collection_efficiency','tariff','capital_efficiency'):
                d['toggles'][prefix+'_'+lever+'_enabled'] = True
        passes, _, _ = cumulative_passes([d])
        bi = passes[0]['years'].index(d['period']['baseline_year'])
        for sector in ('water_supply','sanitation'):
            references = [p[sector]['scenario_reference_billed_volume_million_m3'] for p in passes]
            self.assertTrue(all(r == references[0] for r in references))
            for p in passes:
                s = p[sector]
                q = s['scenario_connection_revenue']['calibration']['consumption_m3']
                for t in range(bi+1,len(p['years'])):
                    self.assertAlmostEqual(s['scenario_household_billed_volume_million_m3'][t],
                                           q*s['scenario_hh'][0][t-1])
            metric = [np.array(p[sector]['scenario_endline_financing_requirement']) for p in passes]
            np.testing.assert_allclose(sum(metric[i+1]-metric[i] for i in range(len(metric)-1)),
                                       metric[-1]-metric[0], atol=1e-9)

    def test_connection_cash_does_not_silently_expand_debt_eligibility(self):
        d = example()
        d['connection_revenue'] = {'water':configuration(billed_share_basic=0)}
        by = d['period']['baseline_year']
        d['utility_debt'] = {'schema_version':1, 'water':dict(
            enabled=True, allocation_share=.5, annual_real_interest_rate=.05,
            disbursement_year=by+1, maturity_year=d['period']['forecast_end_year'],
            repayment_structure='annuity', principal_grace_years=0, loan_ceiling=1)}
        result = calculate(coerce_to_engine(d))
        sec = result['water_supply']
        self.assertTrue(sec['scenario_utility_debt']['enabled'])
        self.assertGreater(max(sec['scenario_connection_net_cash']),0)
        self.assertEqual(max(sec['scenario_eligible_additional_revenue']),0)
        self.assertEqual(sec['scenario_utility_debt']['accepted_principal'],0)
        self.assertEqual(sec['scenario_hh'],sec['scenario_without_utility_debt_hh'])

    def test_exports_and_national_intensive_rates(self):
        from export_data import scenario_xlsx, per_year_table
        from openpyxl import load_workbook
        from deck_aggregate import aggregate
        from revenue_export import revenue_table, append_revenue_slides
        from pptx import Presentation
        d = example()
        d['connection_revenue'] = {s: configuration() for s in ('water','sanitation')}
        result = calculate(coerce_to_engine(d))
        h, rows = per_year_table(result, d, 'water_supply')
        self.assertEqual(len(h),len(rows[-1]))
        self.assertTrue(any('Connection net cash' in v for v in h))
        wb = load_workbook(scenario_xlsx(d))
        self.assertIn('Water — revenue details', wb.sheetnames)
        self.assertIn('Revenue assumptions', wb.sheetnames)
        headers, rows = revenue_table(result,'water_supply','LCU')
        self.assertEqual(len(headers),len(rows[-1]))
        agg = aggregate([result,result])
        for prefix in ('','scenario_'):
            self.assertEqual(agg['water_supply'][prefix+'applicable_tariff'],result['water_supply'][prefix+'applicable_tariff'])
            self.assertEqual(agg['water_supply'][prefix+'applicable_collection_ratio'],result['water_supply'][prefix+'applicable_collection_ratio'])
        mixed = aggregate([result,calculate(coerce_to_engine(example()))])
        self.assertTrue(mixed['water_supply']['connection_revenue']['mixed'])
        prs = Presentation()
        append_revenue_slides(prs, {'urban':result}, 'LCU')
        stream = io.BytesIO()
        prs.save(stream)
        reopened = Presentation(io.BytesIO(stream.getvalue()))
        self.assertTrue(any('revenue details' in s.shapes[0].text for s in reopened.slides))
        from export_deck import build_deck
        branded = Presentation(build_deck({'urban':d, 'rural':d}))
        revenue_slides = [s for s in branded.slides
                          if any('revenue details' in sh.text for sh in s.shapes if sh.has_text_frame)]
        self.assertTrue(revenue_slides)
        for slide in revenue_slides:
            self.assertIsNotNone(slide.notes_slide.notes_text_frame)
            self.assertIn('annual_rows',slide.notes_slide.notes_text_frame.text)


if __name__ == '__main__':
    unittest.main()
