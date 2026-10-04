import unittest

from openpyxl import load_workbook
from pptx import Presentation

from currency_export import validate_currency_display
from export_data import _currency_table, chart_xlsx, scenario_csv, scenario_xlsx, table_xlsx
from export_deck import build_deck
from export_pptx import create_pptx
from demo_adapter import coerce_to_engine
from model.engine import calculate
from test_support import frontend_defaults


class CurrencyExportTests(unittest.TestCase):
    def test_fixed_rate_validation_and_usd_identity(self):
        settings = validate_currency_display({
            'mode': 'usd', 'sourceCurrency': 'NPR', 'localPerUsd': 132.5,
            'rateReferenceYear': 2025, 'sourceNote': 'Central bank annual average',
        }, ['npr'])
        self.assertAlmostEqual(settings['factor'], 1 / 132.5)
        self.assertIn('2025 reference rate', settings['rate_note'])
        self.assertEqual(settings['source_note'], 'Central bank annual average')

        identity = validate_currency_display({'mode': 'usd'}, ['USD'])
        self.assertEqual(identity['factor'], 1)
        self.assertEqual(identity['rate'], 1)

        with self.assertRaisesRegex(ValueError, 'finite and greater than zero'):
            validate_currency_display({
                'mode': 'usd', 'sourceCurrency': 'NPR', 'localPerUsd': 0,
                'rateReferenceYear': 2025,
            }, ['NPR'])
        with self.assertRaisesRegex(ValueError, 'same source currency'):
            validate_currency_display({
                'mode': 'usd', 'sourceCurrency': 'NPR', 'localPerUsd': 132.5,
                'rateReferenceYear': 2025,
            }, ['NPR', 'USD'])

    def test_only_explicit_money_columns_are_converted(self):
        headers = ['Year', 'Total HH (M)', 'Need (npr M)', 'Stock (NPR B)']
        rows = [[2040, 2.5, 1325, 2.65]]
        display = {'mode': 'usd', 'source_currency': 'NPR', 'display_currency': 'USD', 'factor': 1 / 132.5}
        out_headers, out_rows, money_columns = _currency_table(headers, rows, display)
        self.assertEqual(money_columns, [2, 3])
        self.assertEqual(out_headers[1], 'Total HH (M)')
        self.assertEqual(out_headers[2], 'Need (US$ M)')
        self.assertAlmostEqual(out_rows[0][2], 10)
        self.assertAlmostEqual(out_rows[0][3], 0.02)

    def test_per_table_and_chart_workbooks_include_currency_metadata(self):
        display = validate_currency_display({
            'mode': 'usd', 'sourceCurrency': 'NPR', 'localPerUsd': 100,
            'rateReferenceYear': 2025,
        }, ['NPR'])
        sheets = [{'name': 'Data', 'headers': ['Year', 'Value (US$ M)'], 'rows': [[2040, 3.2]]}]
        for output in (
            table_xlsx(sheets, currency_display=display),
            chart_xlsx('Data chart', sheets, currency_display=display),
        ):
            workbook = load_workbook(output, data_only=True)
            self.assertIn('Export metadata', workbook.sheetnames)
            self.assertEqual(workbook['Export metadata']['B2'].value, 'USD')
            self.assertEqual(workbook['Export metadata']['B6'].value, 100)

    def test_scenario_exports_keep_source_detail_and_rate_metadata(self):
        inputs = frontend_defaults()
        display = validate_currency_display({
            'mode': 'usd', 'sourceCurrency': 'NPR', 'localPerUsd': 100,
            'rateReferenceYear': 2025, 'sourceNote': 'Example fixed rate',
        }, ['NPR'])

        csv = scenario_csv(inputs, currency_display=display)
        self.assertIn('US$ M', csv)
        self.assertIn('source-currency monetary detail', csv)
        self.assertIn('2025 reference rate', csv)

        workbook = load_workbook(scenario_xlsx(inputs, currency_display=display), data_only=True)
        self.assertIn('Export metadata', workbook.sheetnames)
        self.assertIn('Water — forecast (USD)', workbook.sheetnames)
        self.assertIn('Water — local detail', workbook.sheetnames)
        metadata = workbook['Export metadata']
        self.assertEqual(metadata['B2'].value, 'USD')
        self.assertEqual(metadata['B6'].value, 100)

        usd = workbook['Water — forecast (USD)']
        local = workbook['Water — local detail']
        checked = False
        usd_headers = {cell.value: cell.column for cell in usd[1]}
        for local_column in range(2, local.max_column + 1):
            local_header = local.cell(1, local_column).value
            usd_header = str(local_header).replace('(NPR M)', '(US$ M)').replace('(NPR B)', '(US$ B)')
            usd_column = usd_headers.get(usd_header)
            if not usd_column:
                continue
            for row_number in range(2, min(usd.max_row, local.max_row) + 1):
                local_value = local.cell(row_number, local_column).value
                usd_value = usd.cell(row_number, usd_column).value
                if isinstance(local_value, (int, float)) and local_value and isinstance(usd_value, (int, float)):
                    self.assertAlmostEqual(usd_value, local_value / 100, places=5)
                    checked = True
                    break
            if checked:
                break
        self.assertTrue(checked, 'expected a matching nonzero local/USD monetary result')

    def test_both_powerpoint_exports_include_usd_presentation_notes(self):
        inputs = frontend_defaults()
        display = validate_currency_display({
            'mode': 'usd', 'sourceCurrency': 'NPR', 'localPerUsd': 100,
            'rateReferenceYear': 2025, 'sourceNote': 'Example fixed rate',
        }, ['NPR'])

        live_result = calculate(coerce_to_engine(inputs))
        result_deck = Presentation(create_pptx(live_result, inputs, currency_display=display))
        result_text = '\n'.join(shape.text for slide in result_deck.slides for shape in slide.shapes
                                if getattr(shape, 'has_text_frame', False))
        self.assertIn('US$1 = 100 NPR', result_text)
        self.assertIn('Constant-price model values translated', result_text)

        branded_deck = Presentation(build_deck({'urban': inputs}, currency_display=display))
        cover_text = '\n'.join(shape.text for shape in branded_deck.slides[0].shapes
                               if getattr(shape, 'has_text_frame', False))
        self.assertIn('US$1 = 100 NPR', cover_text)
        investment_tables = [
            shape.table
            for slide in branded_deck.slides
            for shape in slide.shapes
            if getattr(shape, 'has_table', False)
            and shape.table.rows
            and shape.table.cell(0, 0).text == 'Urban water'
        ]
        self.assertTrue(investment_tables)
        headers = [cell.text for cell in investment_tables[0].rows[0].cells]
        self.assertTrue(all('USD b' in header for header in headers[1:]))
        self.assertFalse(any('US$ m' in header for header in headers))


if __name__ == '__main__':
    unittest.main()