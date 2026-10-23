"""Guard the patched dependencies and the integrations that use them."""

from importlib.metadata import version
from io import BytesIO
from pathlib import Path
import unittest

import h11
from PIL import Image
from pptx import Presentation
from pptx.util import Inches


class DependencySecurityTests(unittest.TestCase):
    def test_security_pins_are_installed_and_above_reported_fixed_versions(self):
        # These floors cover all reported advisory ranges, not just the
        # advisories an individual audit service happens to return.
        floors = {
            "pillow": (13, 0, 0),
            "anyio": (4, 14, 2),
            "h11": (0, 16, 0),
            "idna": (3, 15),
        }
        pins = {}
        for line in Path(__file__).with_name("requirements.txt").read_text().splitlines():
            if "==" in line and not line.startswith("#"):
                name, pinned = line.strip().split("==")
                pins[name.lower()] = pinned
        for name, floor in floors.items():
            with self.subTest(package=name):
                self.assertIn(name, pins)
                self.assertGreaterEqual(tuple(map(int, pins[name].split("."))), floor)
                self.assertEqual(version(name), pins[name])

    def test_h11_rejects_malformed_chunk_terminator(self):
        connection = h11.Connection(h11.SERVER)
        connection.receive_data(
            b"POST / HTTP/1.1\r\nHost: example.test\r\n"
            b"Transfer-Encoding: chunked\r\n\r\n5\r\nAAAAAXX0\r\n\r\n"
        )
        self.assertIsInstance(connection.next_event(), h11.Request)
        self.assertIsInstance(connection.next_event(), h11.Data)
        with self.assertRaises(h11.RemoteProtocolError):
            connection.next_event()

    def test_pillow_image_roundtrips_through_powerpoint(self):
        png = BytesIO()
        Image.new("RGB", (16, 16), "blue").save(png, format="PNG")
        png.seek(0)
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_picture(png, Inches(1), Inches(1))
        output = BytesIO()
        deck.save(output)
        output.seek(0)
        picture = Presentation(output).slides[0].shapes[0]
        with Image.open(BytesIO(picture.image.blob)) as image:
            self.assertEqual(image.size, (16, 16))
            self.assertEqual(image.getpixel((0, 0)), (0, 0, 255))


if __name__ == "__main__":
    unittest.main()
