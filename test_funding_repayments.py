"""Full injection, fixed repayment and addition-only revenue acceptance contracts."""
import copy
import unittest
import numpy as np
from model.fixed_loan import fixed_schedule
from model.utility_debt import indicative_principal
from model.revenue_reconciliation import reconcile_revenue
from model.engine import calculate
from demo_adapter import coerce_to_engine
from deck_aggregate import _aggregate_utility_debt
from loan_reporting import loan_tables
from test_consolidated_revenue import project
import test_indicative_loans as loan_fixture
from test_support import frontend_defaults


class FundingRepaymentTests(unittest.TestCase):
    def test_numeric_annuity_and_small_rate(self):
        for rate in (0, 1e-12, .05):
            p, _ = indicative_principal(20, rate, 10)
            plan = fixed_schedule(p, 20, rate, 10, 2028, range(2025,2041))
            schedule = plan['schedule']
            self.assertEqual(schedule[0]['debt_service'], 0)
            self.assertEqual(schedule[1]['year'], 2029)
            self.assertEqual(schedule[-1]['year'], 2038)
            self.assertEqual(schedule[-1]['closing_principal'], 0)
            self.assertAlmostEqual(sum(row['principal_payment'] for row in schedule),p)
            for row in schedule[1:]:
                self.assertAlmostEqual(row['principal_payment']+row['interest_payment'],20)
                self.assertAlmostEqual(row['interest_payment'], row['opening_principal']*rate)
            self.assertEqual(plan['debt_service'][-2:],[0,0])
            if rate == .05:
                self.assertAlmostEqual(p,154.434698584,places=8)
                self.assertAlmostEqual(schedule[1]['interest_payment'],7.721734929,places=8)
                self.assertAlmostEqual(schedule[1]['closing_principal'],142.156433513,places=8)

    def test_no_payments_zero_loan(self):
        for years in ([2025,2026], [2025,2030]):
            plan = fixed_schedule(0,0,None,None,None,years)
            self.assertEqual(plan['schedule'],[])
            self.assertEqual(sum(plan['debt_service']),0)

    def test_full_one_time_independent_of_factors(self):
        for share, execution in ((.2,.8),(0,.8),(.2,0),(1,1)):
            args = dict(capex_pct=share,execution_rate=execution,
                        financial_settings=dict(injection_amount=100,injection_start_year=2028,
                                                injection_end_year=2028,injection_mode='one_time'))
            sec = project(n=5,injection_enabled=True,**args)
            self.assertEqual(sec['exogenous_injection_cash'],[0,0,0,100,0])
            off = project(n=5,injection_enabled=False,**args)
            self.assertEqual(sum(off['exogenous_injection_cash']),0)
            args['financial_settings']['injection_mode']='recurring'
            recurring = project(n=5,injection_enabled=True,**args)
            self.assertAlmostEqual(recurring['exogenous_injection_cash'][3],100*share*execution)

    def test_both_service_injections(self):
        d = frontend_defaults()
        for sector,prefix,key in (('water','ws','water_supply'),('sanitation','san','sanitation')):
            d['toggles'][prefix+'_exogenous_injection_enabled']=True
            d[sector+'_interventions'].update(fin_injection_amount=100,
                fin_injection_start_year=2028,fin_injection_mode='one_time')
        result = calculate(coerce_to_engine(d))
        for key in ('water_supply','sanitation'):
            values = result[key]['scenario_exogenous_injection_cash']
            self.assertEqual(values[result['years'].index(2028)],100)
            self.assertEqual(sum(values),100)
            self.assertEqual(sum(result[key]['exogenous_injection_cash']),0)

    def test_funding_order_and_restricted_identity(self):
        plan = dict(disbursement=[0,40,0,0],principal_payment=[0,0,20,20],
                    interest_payment=[0,0,0,0],debt_service=[0,0,20,20])
        sec = project(initial=(1,0,0,0,0), full_budget=np.array([0,0,100,10]),utility_debt_execution=plan,
                      cost_sm=700,cost_basic=350)
        for i in range(4):
            self.assertEqual(sec['available_after_debt_service'][i],
                             max(0,sec['available_total'][i])-sec['debt_service_paid'][i])
            self.assertAlmostEqual(sec['utility_debt_cash_opening'][i]+sec['utility_debt_disbursement'][i],
                sec['utility_debt_investment_used'][i]+sec['utility_debt_cash_closing'][i])
        # No servicing sources selected: baseline cannot repay contractual debt.
        self.assertEqual(sec['available_after_debt_service'][2],100)
        self.assertEqual(sec['available_after_debt_service'][3],10)
        self.assertEqual(sec['cash_deficit'][3],20)
        self.assertGreater(sec['utility_debt_cash_closing'][3],0)

    def test_overlap_cannot_consume_baseline_and_identity(self):
        for candidate, eligible, sales, net, volume in ((-10,8,15,0,115),(20,30,30,0,130),
                                                       (20,5,15,15,130),(0,0,0,0,100)):
            positive=max(0,candidate)
            row=reconcile_revenue(100+positive,100,sales,min(positive,eligible),2,.7,3,.9,implementation_cost=5)
            self.assertEqual(row['connection_volume_million_m3'],net)
            self.assertEqual(row['billed_volume_million_m3'],volume)
            self.assertAlmostEqual(row['additional_net_cash'],volume*3*.9-100*2*.7-5)
            if candidate==20 and eligible==5:
                np.testing.assert_allclose([row[k] for k in ('connection_net_cash','nrw_net','collection_cash','tariff_cash')],[21,16,52,117])
        from model.utility_revenue import RevenueInputError
        for raw, overlap in ((90,0),(120,21)):
            with self.assertRaises(RevenueInputError):
                reconcile_revenue(raw,100,30,overlap,2,.7,3,.9)

    def test_horizon_obligations_aggregation_and_export_nulls(self):
        helper=loan_fixture.IndicativeLoanTests()
        early=helper.solve()[0]
        late=helper.solve(disbursement_year=2027,loan_term_years=5,annual_real_interest_rate=0,
                          future_collection=20,future_tariff=4)[0]
        self.assertGreater(early['horizon_closing_principal'],0)
        self.assertEqual(early['remaining_contractual_payments'],9)
        combined=_aggregate_utility_debt([early,late])
        self.assertAlmostEqual(combined['horizon_closing_principal'],
                               early['horizon_closing_principal']+late['horizon_closing_principal'])
        payments={r['year']:r['debt_service'] for r in combined['repayment_schedule']}
        self.assertEqual(payments[2027],early['fixed_annual_debt_service'])
        self.assertEqual(payments[2028],early['fixed_annual_debt_service']+late['fixed_annual_debt_service'])
        headers,rows=loan_tables(combined,'USD',.5)[1]
        final=rows[-1]
        self.assertIsNone(final[headers.index('Ordinary funds after debt service (USD M)')])
        self.assertEqual(final[headers.index('Closing outstanding principal (USD M)')],0)
        self.assertFalse(any('Deferred' in str(cell) for row in rows for cell in row))

    def test_current_csv_excel_and_pptx_repayment_tables(self):
        import io
        from openpyxl import load_workbook
        from pptx import Presentation
        from export_data import scenario_csv,scenario_xlsx
        from export_deck import build_deck
        d=frontend_defaults()
        d['toggles']['ws_collection_efficiency_enabled']=True
        d['water_interventions']['ce_target_ratio']=.99
        d['utility_debt']['water']=dict(enabled=True,allocation_share=.2,
            annual_real_interest_rate=.05,loan_term_years=10,disbursement_year=2028,
            revenue_sources=['collection'])
        csv=scenario_csv(copy.deepcopy(d))
        self.assertIn('Scheduled principal',csv)
        self.assertIn('Ordinary funds after debt service',csv)
        self.assertNotIn('repayments not modeled',csv)
        workbook=load_workbook(scenario_xlsx(copy.deepcopy(d)),data_only=True)
        self.assertTrue(any('Scheduled debt service' in str(value) for sheet in workbook
                            for row in sheet.values for value in row))
        import zipfile
        asset=build_deck({'urban':copy.deepcopy(d)})
        with zipfile.ZipFile(asset) as archive:
            self.assertEqual(len(archive.namelist()),len(set(archive.namelist())))
        asset.seek(0)
        deck=Presentation(asset)
        tables=[shape.table for slide in deck.slides for shape in slide.shapes if shape.has_table]
        self.assertTrue(any('Scheduled principal' in cell.text for table in tables
                            for row in table.rows for cell in row.cells))


if __name__ == '__main__':
    unittest.main()
