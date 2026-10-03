"""Reporting/export parity and safe, explicit saved-scenario compatibility."""
import copy
import csv
import io
import json
import unittest
import warnings
import zipfile

import numpy as np
from openpyxl import load_workbook
from pptx import Presentation

from demo_adapter import frontend_defaults, coerce_to_engine, financial_toggles
from input_compatibility import migrate_input_compatibility
from model.engine import calculate
from reporting import annual_reporting_tables, assumption_rows
from export_data import scenario_csv, scenario_xlsx
from export_pptx import create_pptx
from export_deck import build_deck


class ReportingCompatibilityTests(unittest.TestCase):
    def test_old_financing_is_not_activated_and_history_is_preserved(self):
        old = frontend_defaults()
        old.pop('reporting_schema_version')
        old['water_interventions'].pop('borrow_rate_basis')
        old['water_interventions'].pop('borrow_contract_principal')
        old['water_interventions']['cash_allocation_alpha'] = .7
        old['toggles']['ws_borrowing_enabled'] = True
        before = copy.deepcopy(old)
        migrated = migrate_input_compatibility(old)
        self.assertEqual(old, before)
        self.assertFalse(migrated['toggles']['ws_borrowing_enabled'])
        self.assertFalse(financial_toggles(old)['ws_borrowing_enabled'])
        self.assertEqual(migrated['water_interventions']['cash_allocation_alpha'], 0)
        self.assertEqual(migrated['legacy_financing_settings']['water_interventions']['cash_allocation_alpha'], .7)
        for key in ('period', 'water_service', 'sanitation_service', 'macro', 'population', 'custom_interventions'):
            self.assertEqual(migrated[key], before[key])
        self.assertTrue(any('interest basis was not recorded' in n for n in migrated['migration_notes']))
        self.assertEqual(migrate_input_compatibility(json.loads(json.dumps(migrated))), migrated)
        result = calculate(coerce_to_engine(old))
        self.assertEqual(result['water_supply']['scenario_loan_principal'], 0)

    def test_supported_saves_preserve_explicit_financing_choices(self):
        current = frontend_defaults()
        current['toggles']['ws_borrowing_enabled'] = True
        current['water_interventions'].update(cash_allocation_alpha=.6, borrow_rate_basis='nominal',
                                               borrow_contract_principal=123)
        for remove_version in (False, True):
            inputs = copy.deepcopy(current)
            if remove_version:
                inputs.pop('reporting_schema_version')
            migrated = migrate_input_compatibility(inputs)
            self.assertEqual(migrated['water_interventions'], inputs['water_interventions'])
            self.assertTrue(migrated['toggles']['ws_borrowing_enabled'])

    def inputs(self):
        inputs = frontend_defaults()
        inputs['toggles'].update(ws_tariff_enabled=True, ws_borrowing_enabled=True)
        inputs['water_interventions'].update(cash_allocation_alpha=.5, borrow_term_years=25)
        return inputs

    def test_csv_and_xlsx_share_full_precision_annual_engine_rows(self):
        inputs = self.inputs()
        result = calculate(coerce_to_engine(inputs))
        parsed = list(csv.reader(io.StringIO(scenario_csv(inputs))))
        wb = load_workbook(scenario_xlsx(inputs), data_only=True)
        for sk, name in (('water_supply', 'Water'), ('sanitation', 'Sanitation')):
            for i, table in enumerate(annual_reporting_tables(result, inputs, sk), 1):
                csv_label = 'WATER SUPPLY' if sk == 'water_supply' else 'SANITATION'
                csv_start = parsed.index([csv_label + ' — ' + table['title']])
                self.assertEqual(parsed[csv_start+1], table['headers'])
                csv_rows = parsed[csv_start+2:csv_start+2+len(table['rows'])]
                np.testing.assert_allclose(np.asarray(csv_rows, float), table['rows'], rtol=1e-13, atol=1e-13)
                sheet = wb[f'{name} annual {i}']
                self.assertEqual(list(next(sheet.values)), table['headers'])
                np.testing.assert_allclose(list(sheet.values)[1:], table['rows'], rtol=1e-13, atol=1e-13)
        self.assertIn('Methodology and migration', wb.sheetnames)

    def test_both_ppt_paths_include_complete_reporting_and_unique_zip_parts(self):
        inputs = self.inputs()
        result = calculate(coerce_to_engine(inputs))
        with warnings.catch_warnings():
            warnings.filterwarnings('error', message='Duplicate name:.*')
            for deck in (create_pptx(result, inputs), build_deck({'urban': inputs, 'rural': inputs})):
                zipped = zipfile.ZipFile(deck)
                self.assertEqual(len(zipped.namelist()), len(set(zipped.namelist())))
                self.assertIsNone(zipped.testzip())
                prs = Presentation(deck)
                text = '\n'.join(shape.text for slide in prs.slides for shape in slide.shapes if shape.has_text_frame)
                for title in ('Annual investment requirements', 'Annual financing sources',
                              'Annual and cumulative financing gaps', 'Operating cash and debt',
                              'Basic coverage and gaps', 'Residual additional public financing'):
                    self.assertIn(title, text)
                tables = [shape.table for slide in prs.slides for shape in slide.shapes if shape.has_table]
                value = result['water_supply']['scenario_additional_net_utility_cash'][-1]
                self.assertTrue(any(cell.text == f'{value:,.6f}' for table in tables for row in table.rows for cell in row.cells))

    def test_migration_assumptions_survive_export_and_review_acknowledgment(self):
        inputs = frontend_defaults()
        inputs.pop('reporting_schema_version')
        inputs['water_interventions'].pop('borrow_rate_basis')
        migrated = migrate_input_compatibility(inputs)
        migrated['migration_reviewed'] = True
        notes = [row[1] for row in assumption_rows(migrated)]
        self.assertTrue(any('interest basis was not recorded' in note for note in notes))
        self.assertIn('interest basis was not recorded', scenario_csv(migrated))


if __name__ == '__main__':
    unittest.main()