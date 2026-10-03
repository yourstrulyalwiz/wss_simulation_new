"""Deterministic residual-ledger acceptance cases, in actual dollars/households."""

import unittest
from types import SimpleNamespace

import numpy as np
from openpyxl import load_workbook
from pptx import Presentation

from demo_adapter import coerce_to_engine, frontend_defaults
from export_data import per_year_table, scenario_csv, scenario_xlsx
from export_deck import build_deck
from model.engine import calculate
from model.gap_attribution import attribute_gap
from model.water_supply import sector_bau


def fixture(available, replacement, *, sanitation=False, years=2, fully_sm_target=False):
    # 100 SM + 200 Basic; target 200 SM + 100 Basic: exactly 100 upgrades.
    # All amounts are converted to the engine's millions convention.
    ctx = {
        'n': years, 'bi': 0, 'years': np.arange(2025, 2025 + years),
        'total_hh': np.full(years, 300 / 1e6),
        'population': np.full(years, 900 / 1e6),
        'forecast_flag': np.array([0.] + [1.] * (years - 1)),
        'perf_flag': np.array([0.] + [1.] * (years - 1)),
        'end_asis_year': 2025, 'gdp_real_local': np.ones(years),
    }
    return sector_bau(
        ctx, period=SimpleNamespace(model_start_year=2025, baseline_year=2025,
                                    target1_year=2026, target2_year=2026),
        pct_start=[1/3, 2/3, 0, 0, 0], pct_base=[1/3, 2/3, 0, 0, 0],
        tgt1=[1, 0, 0, 0, 0] if fully_sm_target else [2/3, 1/3, 0, 0, 0],
        tgt2=[1, 0, 0, 0, 0] if fully_sm_target else [2/3, 1/3, 0, 0, 0],
        cost_sm=1000, cost_basic=1000,
        full_budget=np.zeros(years), capex_pct=1, growth_capex_pct=0,
        planned_list=[], nonhh_pct=0, asset_life=300000/replacement if replacement else float('inf'),
        capex_adder=0, hist_all_proportional=not sanitation,
        target_adjusted=sanitation, basic_share=0,
        extra_cash=np.array([0.] + [available / 1e6] * (years - 1)),
    )


class ResidualFinancingGapTests(unittest.TestCase):
    def test_five_deterministic_acceptance_cases_for_both_sector_modes(self):
        cases = [(40000, 0, 40, 60000), (40000, 10000, 30, 70000),
                 (5000, 10000, 0, 105000), (0, 10000, 0, 110000),
                 (-5000, 10000, 0, 115000)]
        for sanitation in (False, True):
            for available, replacement, upgrades, gap in cases:
                with self.subTest(sanitation=sanitation, available=available, replacement=replacement):
                    sec = fixture(available, replacement, sanitation=sanitation)
                    self.assertAlmostEqual((sec['bau_hh'][0][1] - sec['bau_hh'][0][0]) * 1e6, upgrades)
                    self.assertAlmostEqual(sec['financing_gap'][1] * 1e6, gap)
                    self.assertAlmostEqual(sec['replacement_capex'][1] * 1e6, replacement)
                    self.assertAlmostEqual(sec['bau_replacement_capex'][1] * 1e6, replacement)
                    self.assertEqual(sec['financing_gap'][0], 0)
                    self.assertAlmostEqual(sum(r[1] for r in sec['financing_gap_by_service']) * 1e6, gap)
                    self.assertAlmostEqual(sec['connection_purchase_capital'][1] * 1e6, upgrades * 1000)

    def test_unequal_stock_bases_reserve_and_credit_are_capped_separately(self):
        for available in (0, 5000, 40000, 500000):
            sec = fixture(available, 10000, years=4)
            for i in (2, 3):
                reserve = min(max(available / 1e6, 0), max(sec['bau_replacement_capex'][i], 0))
                credit = min(reserve, max(sec['replacement_capex'][i], 0))
                self.assertAlmostEqual(sec['replacement_reserved'][i], reserve)
                self.assertAlmostEqual(sec['replacement_credit'][i], credit)
                self.assertAlmostEqual(sec['financing_gap'][i],
                    sec['new_capex_total'][i] + max(sec['replacement_capex'][i] - credit, 0))
            self.assertNotEqual(sec['bau_replacement_capex'][2], sec['replacement_capex'][2])
        # With surplus cash the coverage-stock reserve exceeds the reported need.
        self.assertGreater(sec['replacement_reserved'][2], sec['replacement_credit'][2])

    def test_fully_funded_target_does_not_recredit_excess_capital(self):
        for sanitation in (False, True):
            sec = fixture(500000, 10000, sanitation=sanitation, fully_sm_target=True)
            # The existing zero-target rung uses a tiny CAGR endpoint epsilon.
            self.assertAlmostEqual(sec['new_capex_total'][1], 0, places=9)
            self.assertAlmostEqual(sec['financing_gap'][1], 0, places=9)
            self.assertAlmostEqual(sec['connection_purchase_capital'][1], 0.2)
            self.assertAlmostEqual(sec['unallocated_positive_capital'][1], 0.29)
            self.assertAlmostEqual(sec['replacement_credit'][1], 0.01)

    def test_service_credit_and_deficit_attribution(self):
        sm, basic, sm_credit, basic_credit = attribute_gap(60, 40, 30, 10, 20, 0.9, 14)
        self.assertEqual((sm_credit, basic_credit), (15, 5))
        self.assertAlmostEqual(sm, 84)
        self.assertAlmostEqual(basic, 50)
        self.assertAlmostEqual(sm + basic, 134)
        self.assertEqual(attribute_gap(0, 0, 0, 0, 0, 0.4, 10), (6, 4, 0, 0))

    def test_engine_csv_excel_and_powerpoint_use_same_residual_ledger(self):
        inputs = frontend_defaults()
        result = calculate(coerce_to_engine(inputs))
        wb = load_workbook(scenario_xlsx(inputs), data_only=True)
        csv_data = scenario_csv(inputs)
        if isinstance(csv_data, bytes):
            csv_data = csv_data.decode()
        self.assertIn('Replacement credit', csv_data)
        self.assertIn('Cash deficit', csv_data)
        for sector in ('water_supply', 'sanitation'):
            headers, rows = per_year_table(result, inputs, sector)
            for label, key in [('Replacement credit', 'replacement_credit'), ('Cash deficit', 'cash_deficit')]:
                col = headers.index(f'{label} — BAU (NPR M)')
                self.assertAlmostEqual(rows[-1][col], result[sector][key][-1], delta=0.000051)
            ws = wb[f"{'Water' if sector == 'water_supply' else 'Sanitation'} — forecast"]
            self.assertEqual(list(next(ws.values)), headers)
            self.assertEqual(list(ws.values)[-1], tuple(rows[-1]))
        deck = Presentation(build_deck({'urban': inputs}))
        investment_tables = [
            sh.table for slide in deck.slides for sh in slide.shapes
            if sh.has_table and any('Remaining financing gap' in cell.text
                                   for row in sh.table.rows for cell in row.cells)]
        self.assertGreaterEqual(len(investment_tables), 2)
        for table in investment_tables:
            labels = [row.cells[0].text for row in table.rows]
            self.assertIn('Replacement credit (D)', labels)
            self.assertIn('Cash deficit (E)', labels)
            # Values are displayed in billions, to one decimal.
            a, b, d, e, gap = [float(table.cell(labels.index(label), len(table.columns)-1).text.replace(',', ''))
                              for label in ('Residual new-service cost (A)', 'Replacement capex needed (B)',
                                            'Replacement credit (D)', 'Cash deficit (E)',
                                            'Remaining financing gap (C − D + E)')]
            self.assertAlmostEqual(gap, a + b - d + e, delta=0.21)
        texts = '\n'.join(sh.text for slide in deck.slides for sh in slide.shapes if sh.has_text_frame)
        self.assertNotIn('current spending covers', texts)
        self.assertNotIn('a gap of about', texts)


if __name__ == '__main__':
    unittest.main()