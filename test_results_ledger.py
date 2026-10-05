import unittest
import numpy as np
from openpyxl import load_workbook
from export_data import table_xlsx
from model.engine import calculate
from demo_adapter import coerce_to_engine
from test_utility_revenue import example
from test_residual_financing_gap import fixture


class ResultsLedgerTests(unittest.TestCase):
    def assert_report(self, sec, prefix=''):
        def series(key):
            return np.asarray(sec[prefix + key])
        identities = [
            ('annual_planned_expansion_cost_by_service', 'annual_planned_expansion_cost'),
            ('closing_outstanding_expansion_by_service', 'closing_outstanding_expansion'),
            ('sector_funded_expansion_by_service', 'sector_funded_expansion'),
            ('externally_funded_expansion_by_service', 'externally_funded_expansion'),
            ('replacement_funding_applied_by_service', 'replacement_reserved'),
            ('endline_financing_requirement_by_service', 'endline_financing_requirement'),
        ]
        for service_key, total_key in identities:
            np.testing.assert_allclose(series(service_key).sum(axis=0), series(total_key), atol=1e-8)
        np.testing.assert_allclose(
            series('prefunding_expansion_cost_by_service').sum(axis=0) + series('replacement_capex'),
            series('catch_up_requirement'), atol=1e-8)
        np.testing.assert_allclose(
            series('closing_outstanding_expansion_by_service') + series('accumulated_shortfalls_by_service'),
            series('endline_financing_requirement_by_service'), atol=1e-8)
        np.testing.assert_allclose(
            series('new_capex_by_service') + series('replacement_by_service') -
            series('funded_by_service') + series('cash_deficit_by_service'),
            series('financing_gap_by_service'), atol=1e-8)

    def test_reporting_reconciles_under_deficits_surplus_and_both_sector_modes(self):
        for sanitation in (False, True):
            for available in (-5000, 0, 5000, 40000, 500000):
                for replacement in (0, 10000):
                    with self.subTest(sanitation=sanitation, available=available, replacement=replacement):
                        self.assert_report(fixture(available, replacement, sanitation=sanitation, years=4))

    def test_bau_and_scenario_expose_real_service_snapshots(self):
        data = example()
        data['toggles']['ws_collection_efficiency_enabled'] = True
        data['toggles']['ws_tariff_enabled'] = True
        data['water_interventions']['basic_share'] = .4
        data['sanitation_interventions']['basic_share'] = .6
        result = calculate(coerce_to_engine(data))
        for sector in ('water_supply', 'sanitation'):
            self.assert_report(result[sector])
            self.assert_report(result[sector], 'scenario_')

    def test_excel_has_year_columns_unrounded_values_and_frozen_labels(self):
        wb = load_workbook(table_xlsx([{
            'name': 'Water ledger', 'headers': ['Row', 'Unit', 2026, 2027],
            'rows': [['BAU', 'B USD', 1.23456789, None], ['Combined scenario', 'B USD', 2.34567891, 3]],
            'freeze_columns': 2,
        }]))
        ws = wb['Water ledger']
        self.assertEqual(ws.freeze_panes, 'C2')
        self.assertEqual(ws['C1'].value, '2026')
        self.assertEqual(ws['C2'].value, 1.23456789)
        self.assertIsNone(ws['D2'].value)
        self.assertEqual(ws['A3'].value, 'Combined scenario')


if __name__ == '__main__':
    unittest.main()
