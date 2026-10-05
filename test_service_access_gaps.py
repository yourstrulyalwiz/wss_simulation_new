import copy
import json
import unittest
from zipfile import ZipFile
from pathlib import Path
from types import SimpleNamespace

import numpy as np
from openpyxl import load_workbook
from pptx import Presentation

from model.service_gaps import assess_service_gaps, HOUSEHOLD_TOLERANCE, SERVICE_GAP_FIELDS
from model.water_supply import sector_bau
from model.engine import calculate
from demo_adapter import coerce_to_engine
from deck_aggregate import aggregate
from export_data import per_year_table, scenario_csv, scenario_xlsx
from export_deck import build_deck
from export_pptx import create_pptx
from currency_export import validate_currency_display


def diagnostics(sm, basic, ts=70, tb=30, total=100):
    # Real household counts converted to the engine's millions convention.
    return assess_service_gaps([total / 1e6],
                              np.array([[sm], [basic], [total-sm-basic], [0], [0]]) / 1e6,
                              np.array([[ts], [tb], [total-ts-tb], [0], [0]]) / 1e6,
                              [2026])


def projection(sm, basic, *, cash=(0, 0), targets=(70, 30), sanitation=False, replacement=0, basic_share=0):
    n = len(cash)
    total = 100 / 1e6
    ctx = {
        'n': n, 'bi': 0, 'years': np.arange(2025, 2025+n),
        'total_hh': np.full(n, total), 'population': np.full(n, 3*total),
        'forecast_flag': np.array([0.] + [1.]*(n-1)),
        'perf_flag': np.array([0.] + [1.]*(n-1)),
        'end_asis_year': 2025, 'gdp_real_local': np.ones(n),
    }
    initial = np.array([sm, basic, 100-sm-basic, 0, 0]) / 100
    target = np.array([*targets, 100-sum(targets), 0, 0]) / 100
    return sector_bau(
        ctx, period=SimpleNamespace(model_start_year=2025, baseline_year=2025,
                                    target1_year=2026, target2_year=2026),
        pct_start=initial, pct_base=initial, tgt1=target, tgt2=target,
        cost_sm=100, cost_basic=50,
        full_budget=np.zeros(n), capex_pct=1, growth_capex_pct=0,
        planned_list=[], nonhh_pct=0,
        asset_life=(sm*100+basic*50)/replacement if replacement else float('inf'),
        capex_adder=0, hist_all_proportional=not sanitation, target_adjusted=True,
        basic_share=basic_share, extra_cash=np.asarray(cash) / 1e6,
    )


class ServiceAccessGapTests(unittest.TestCase):
    def test_acceptance_matrix_and_partial_target(self):
        # Effective target, basic-only diagnostic, SM gap, combined gap.
        for sm, basic, expected in (
            (80, 20, (20, 0, 0, 0)), (80, 15, (20, 5, 0, 5)),
            (60, 40, (30, 0, 10, 0)), (60, 30, (30, 0, 10, 10)),
            (100, 0, (0, 0, 0, 0)),
        ):
            result = diagnostics(sm, basic)
            actual = [result[field][0]*1e6 for field in (
                'effective_basic_only_target', 'adjusted_basic_only_gap',
                'sm_access_gap', 'at_least_basic_access_gap')]
            np.testing.assert_allclose(actual, expected, atol=.01)
        result = diagnostics(70, 20, 60, 30)
        self.assertEqual(result['at_least_basic_access_gap'][0], 0)
        self.assertEqual(result['sm_access_gap'][0], 0)
        zero_basic = diagnostics(100, 0, 100, 0)
        self.assertEqual(zero_basic['effective_basic_only_target'][0], 0)

    def test_actual_ledger_costs_follow_transitions_both_sectors(self):
        for sanitation in (False, True):
            for sm, basic, transitions, cost in (
                (80, 20, (0, 0), 0), (80, 15, (0, 5), 250),
                (60, 40, (10, 0), 1000), (60, 30, (10, 10), 1500),
                (100, 0, (0, 0), 0),
            ):
                with self.subTest(sanitation=sanitation, sm=sm, basic=basic):
                    result = projection(sm, basic, sanitation=sanitation)
                    np.testing.assert_allclose(np.array(result['closing_outstanding_hh'])[:, -1]*1e6,
                                               transitions, atol=.01)
                    self.assertAlmostEqual(result['closing_outstanding_expansion'][-1]*1e6, cost, places=3)

    def test_funded_overachievement_stays_credited_and_assets_remain(self):
        for sanitation in (False, True):
            result = projection(70, 30, cash=(0, 1000, 0), sanitation=sanitation)
            np.testing.assert_allclose(np.array(result['bau_hh'])[:2, 1:]*1e6, [[80,80], [20,20]])
            np.testing.assert_allclose(np.array(result['closing_outstanding_hh'])[:, 1:], 0)
            np.testing.assert_allclose(result['at_least_basic_access_gap'], 0)
            self.assertGreater(result['household_gap_basic'][-1], 0, 'Legacy raw field retains its meaning')
            self.assertEqual(result['adjusted_basic_only_gap'][-1], 0)
            self.assertAlmostEqual(result['funded_asset_stock'][-1]*1e6, 9500)

    def test_later_basic_entries_close_real_deficit_once(self):
        for sanitation in (False, True):
            result = projection(80, 15, cash=(0, 0, 250, 0), sanitation=sanitation, basic_share=1)
            np.testing.assert_allclose(np.array(result['closing_outstanding_hh'])[1]*1e6, [0,5,0,0], atol=.01)
            self.assertAlmostEqual(result['bau_hh'][1][-1]*1e6, 20)
            self.assertEqual(result['closing_outstanding_expansion'][-1], 0)
        # The genuine five-entry balance also closes directly in the shared ledger.
        from model.expansion_ledger import ExpansionLedger
        ledger = ExpansionLedger(80/1e6, 15/1e6, 3)
        for t, entries in ((1, 0), (2, 5/1e6)):
            ledger.step(t, [70/1e6,30/1e6], [100,50], [0,entries], [0,0], [0,0],
                        replacement=0, unpaid_by_service=[0,0], deficit_by_service=[0,0],
                        ancillary_cost=0, spare_capital=0, asset_stock=0,
                        sector_capital=entries*50, external_capital=0)
        self.assertAlmostEqual(ledger.series['closing_outstanding_hh'][1,1]*1e6, 5)
        self.assertEqual(ledger.series['closing_outstanding_expansion'][2], 0)

    def test_replacement_remains_when_all_access_objectives_are_met(self):
        result = projection(80,20, cash=(0,0,0), replacement=100)
        self.assertEqual(result['closing_outstanding_expansion'][-1], 0)
        self.assertAlmostEqual(result['unfunded_replacement'][-1]*1e6, 100)
        self.assertAlmostEqual(result['endline_financing_requirement'][-1]*1e6, 200)

    def test_validation_and_numerical_residue(self):
        exact = diagnostics(70,30)
        self.assertEqual(exact['adjusted_basic_only_gap'][0],0)
        near = diagnostics(80,20-.005)
        self.assertEqual(near['at_least_basic_access_gap'][0],0)
        for args in ((80,30), (-1,20), (80,float('nan')), (80,20,80,30)):
            with self.assertRaises(ValueError):
                diagnostics(*args)
        with self.assertRaisesRegex(ValueError, 'sum to total'):
            assess_service_gaps([100/1e6], np.array([[70],[20],[0],[0],[0]])/1e6,
                                np.array([[70],[30],[0],[0],[0]])/1e6)
        self.assertEqual(HOUSEHOLD_TOLERANCE*1e6, .01)

    def test_growing_households_and_geography_before_aggregation(self):
        growing = assess_service_gaps(np.array([100,120])/1e6,
                                      np.array([[70,80],[30,40],[0,0],[0,0],[0,0]])/1e6,
                                      np.array([[70,84],[30,36],[0,0],[0,0],[0,0]])/1e6,
                                      [2026,2027])
        np.testing.assert_allclose(growing['sm_access_gap']*1e6,[0,4],atol=.01)
        np.testing.assert_allclose(growing['at_least_basic_access_gap'],0)
        def result(total, sm, basic):
            sec = {key:value.tolist() for key,value in diagnostics(sm,basic,total*.7,total*.3,total).items()}
            return {'years':[2026], 'total_hh':[total/1e6], 'population':[3*total/1e6],
                    'water_supply':sec,'sanitation':copy.deepcopy(sec)}
        national = aggregate([result(100,80,20), result(200,120,60)])
        # The first area's SM surplus must not erase the second area's 20 SM deficit.
        self.assertAlmostEqual(national['water_supply']['sm_access_gap'][0]*1e6,20)
        self.assertAlmostEqual(national['water_supply']['at_least_basic_access_gap'][0]*1e6,20)
        self.assertAlmostEqual(national['water_supply']['sm_overachievement'][0]*1e6,10)

    def test_drc_scenario_mapping_counterfactual_and_exports(self):
        path = Path('profiles/DRC_Mock_Simulation.json')
        original = path.read_bytes()
        bundle = json.loads(original)
        inputs = bundle['inputs']
        result = calculate(coerce_to_engine(inputs))
        disabled = copy.deepcopy(inputs)
        disabled['toggles'] = {key:False for key in disabled['toggles']}
        bau = calculate(coerce_to_engine(disabled))
        for sector in ('water_supply','sanitation'):
            for key in SERVICE_GAP_FIELDS:
                self.assertEqual(result[sector][key],bau[sector][key])
                self.assertEqual(bau[sector][key],bau[sector]['scenario_'+key])
            sec=result[sector]
            np.testing.assert_allclose(sec['scenario_closing_outstanding_hh'][1],
                                       sec['scenario_at_least_basic_access_gap'],atol=HOUSEHOLD_TOLERANCE)
            headers,rows=per_year_table(result,inputs,sector)
            index=headers.index('At-least-basic access gap — basic-entry costing — scenario (M HH)')
            self.assertAlmostEqual(rows[-1][index],sec['scenario_at_least_basic_access_gap'][-1],places=6)
        self.assertIn('At-least-basic access gap',scenario_csv(inputs))
        workbook=load_workbook(scenario_xlsx(inputs),data_only=True)
        self.assertTrue(any('At-least-basic access gap' in str(cell.value)
                            for sheet in workbook for cell in sheet[1]))
        display = validate_currency_display({'mode':'usd','sourceCurrency':'CDF','localPerUsd':2309.58,
                                             'rateReferenceYear':2024,'sourceNote':'User-supplied mock rate'},['CDF'])
        converted = load_workbook(scenario_xlsx(inputs,currency_display=display),data_only=True)
        values = []
        local_forecast = next(sheet for sheet in workbook if sheet.title.startswith('Water — forecast'))
        for sheet in (converted['Water — forecast (USD)'], local_forecast):
            column = next(cell.column for cell in sheet[1]
                          if cell.value == 'At-least-basic access gap — basic-entry costing — scenario (M HH)')
            values.append(sheet.cell(len(result['years'])+1,column).value)
        self.assertEqual(values[0],values[1],'Household diagnostics must not receive currency conversion')
        for output in (create_pptx(result,inputs), build_deck({'urban':inputs})):
            with ZipFile(output) as package:
                names=package.namelist()
                self.assertEqual(len(names),len(set(names)),'PowerPoint part names must be unique')
            deck=Presentation(output)
            text=' '.join(shape.text if shape.has_text_frame else
                          ' '.join(cell.text for row in shape.table.rows for cell in row.cells)
                          if shape.has_table else ''
                          for slide in deck.slides for shape in slide.shapes)
            self.assertIn('At-least-basic access gap (basic-entry costing)',text)
            self.assertIn('Basic-only shortfall after SM credit (diagnostic)',text)
        self.assertEqual(original,path.read_bytes())


if __name__ == '__main__':
    unittest.main()
