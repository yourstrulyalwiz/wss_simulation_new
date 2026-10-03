"""Integration regressions: canonical area gaps, report data and signed changes."""
import unittest
import numpy as np
from pptx import Presentation

from deck_data import build
from demo_adapter import frontend_defaults, coerce_to_engine
from model.engine import calculate
from export_deck import build_deck
from export_pptx import _sector_summary
from reporting import annual_reporting_tables
from validation_fixtures import mixed_area_inputs


class ValidationIntegrationTests(unittest.TestCase):
    def test_geographical_gaps_and_deck_data_use_per_area_reconciliations(self):
        inputs = mixed_area_inputs()
        data = build(inputs)
        national = data['results']['national']
        for sector in ('water_supply', 'sanitation'):
            for key in ('service_gap_display', 'scenario_service_gap_display', 'financing_gap',
                        'scenario_financing_gap', 'scenario_cumulative_financing_gap',
                        'scenario_financing_gap_before_additional_public', 'scenario_additional_net_utility_cash'):
                expected = sum(np.asarray(data['results'][area][sector][key]) for area in ('urban', 'rural'))
                np.testing.assert_allclose(national[sector][key], expected)
            for rung in (0, 1):
                sec = national[sector]
                naive = np.maximum(0, np.asarray(sec['target_hh'][rung]) -
                                   np.asarray(sec['scenario_hh'][rung]))
                canonical = np.asarray(sec['scenario_service_gap_display'][rung])
                self.assertTrue(np.any(canonical > naive + 1e-6), 'fixture must detect cross-area cancellation')
                rows = annual_reporting_tables(national, inputs['urban'], sector)[rung]['rows']
                np.testing.assert_allclose([row[5] for row in rows], canonical)
            summary = _sector_summary(national, inputs['urban'], sector)
            self.assertEqual(summary['unmetSm'], national[sector]['scenario_service_gap_display'][0][-1])
            self.assertEqual(summary['unmetBasic'], national[sector]['scenario_service_gap_display'][1][-1])
            block = data['blocks'][('national', sector)]
            for year, gap in zip(block['service_gap']['years'], block['service_gap']['gap']):
                self.assertEqual(gap, national[sector]['service_gap_display'][0][national['years'].index(year)])

    def test_template_annual_tables_preserve_mixed_area_gap_values(self):
        inputs = mixed_area_inputs()
        data = build(inputs)
        prs = Presentation(build_deck(inputs))
        for sector in ('water_supply', 'sanitation'):
            expected = annual_reporting_tables(data['results']['national'], inputs['urban'], sector)
            for table in expected[:2]:
                actual_rows = []
                for slide in prs.slides:
                    title = '\n'.join(shape.text for shape in slide.shapes if shape.has_text_frame)
                    if 'National · ' + sector.replace('_', ' ').title() + ' · ' + table['title'] not in title:
                        continue
                    for shape in slide.shapes:
                        if shape.has_table:
                            actual_rows += [[float(cell.text.replace(',', '')) for cell in row.cells]
                                            for row in list(shape.table.rows)[1:]]
                self.assertEqual(len(actual_rows), len(table['rows']))
                np.testing.assert_allclose(actual_rows, table['rows'], atol=5.1e-7, rtol=0)

    def test_powerpoint_summary_does_not_hide_negative_intervention_coverage_change(self):
        inputs = frontend_defaults()
        inputs['toggles']['ws_borrowing_enabled'] = True
        inputs['water_interventions'].update(cash_allocation_alpha=0, existing_debt_service=1e9)
        result = calculate(coerce_to_engine(inputs))
        expected = result['water_supply']['scenario_hh'][0][-1] - result['water_supply']['bau_hh'][0][-1]
        self.assertLess(expected, 0)
        self.assertAlmostEqual(_sector_summary(result, inputs, 'water_supply')['addHH'], expected)


if __name__ == '__main__':
    unittest.main()