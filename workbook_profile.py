"""Build a data-only area bundle without inheriting Nepal's model assumptions."""
import math
from pathlib import Path
from openpyxl import load_workbook
from demo_adapter import frontend_defaults
from excel_io import parse_template

PREVIEW_PROFILE_NAME = "DRC_Data_Preview_Settings_Pending"


def _empty(value):
    if isinstance(value, bool):
        return False
    if isinstance(value, (int, float)):
        return None
    if isinstance(value, list):
        return []
    if isinstance(value, dict):
        return {key: _empty(item) for key, item in value.items()}
    return value


def import_area(path, area):
    path = Path(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["WSS Inputs"]
    start, baseline, end = map(int, sheet["A2"].value.split(":")[1:])
    supplied = {}
    for row in sheet.iter_rows(min_row=7, values_only=True):
        if row[0]:
            supplied[str(row[0])] = list(row[2:])
    workbook.close()
    seed = _empty(frontend_defaults())
    seed["country_config"].update(country="Congo, Dem. Rep.",
                                  area="DRC data preview — settings pending", currency="CDF")
    seed["period"].update(model_start_year=start, baseline_year=baseline, forecast_end_year=end,
                          as_is_forecast_start=baseline + 1)
    seed["revenue_bases"] = {}
    seed["revenue_legacy"] = {}
    imported, _ = parse_template(path.read_bytes(), seed)
    targets = []
    for offset in range(baseline - start + 1, end - start + 1):
        for section, prefix in (("water_service", "serv"), ("sanitation_service", "sserv")):
            total = sum((imported[section][f"{prefix}{i}_ts"][offset]
                         if offset < len(imported[section][f"{prefix}{i}_ts"]) else 0)
                        for i in range(1, 6))
            if math.isclose(total, 1, abs_tol=.0001):
                targets.append(start + offset)
    targets = sorted(set(targets))
    for index, year in enumerate(targets[:2], 1):
        imported["period"][f"target{index}_year"] = year
        for section, prefix, target_section in (
            ("water_service", "serv", "water_targets"),
            ("sanitation_service", "sserv", "sanitation_targets"),
        ):
            for rung in range(1, 6):
                values = imported[section][f"{prefix}{rung}_ts"]
                imported[target_section][f"target{index}_{prefix}{rung}"] = (
                    values[year - start] if year - start < len(values) else 0)
    # Preserve every supplied keyed numeric cell, including explicit zeros.
    count = 0
    for key, values in supplied.items():
        section, field = key.split(".")
        if section not in imported or field not in imported[section]:
            raise ValueError(f"Unmapped spreadsheet field: {key}")
        for index, value in enumerate(values):
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                if imported[section][field][index] != value:
                    raise ValueError(f"Import changed {key} in {start + index}")
                count += 1
    imported["profile_metadata"] = {
        "status": "data_preview", "area": area, "source_workbook": path.name,
        "imported_cells": count,
        "notice": "Spreadsheet data only. Unit costs, technical assumptions, economic "
                  "projection settings and intervention settings have not been supplied. "
                  "Simulation calculations and exports are unavailable until these are completed.",
    }
    return imported


def build_preview(urban_path, rural_path):
    urban = import_area(urban_path, "urban")
    rural = import_area(rural_path, "rural")
    if urban["period"] != rural["period"]:
        raise ValueError("Urban and rural analysis periods do not match.")
    return {
        "__wss_bundle": 1, "inputs": urban, "altInputs": {"rural": rural},
        "scope": {"scopeMode": "urban_rural", "areaUrban": True, "areaRural": True},
        "profile_metadata": {
            "status": "data_preview",
            "notice": "Urban and rural GDP series differ; values are retained exactly as uploaded. "
                      "Grey calculated rows are not imported as assumptions or overrides.",
        },
    }